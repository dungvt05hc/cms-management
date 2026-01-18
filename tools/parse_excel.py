import argparse
import json
from pathlib import Path
from typing import List, Tuple

import pandas as pd
from openpyxl import load_workbook


def worksheet_to_dataframe(ws) -> pd.DataFrame:
    """Convert a worksheet to DataFrame, expanding merged cells."""
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return pd.DataFrame()

    max_len = max(len(r) for r in rows)
    data = [list(r) + [None] * (max_len - len(r)) for r in rows]

    # Propagate merged cell values into their covered range for easier parsing.
    for merged in ws.merged_cells.ranges:
        value = ws.cell(merged.min_row, merged.min_col).value
        for row in range(merged.min_row, merged.max_row + 1):
            for col in range(merged.min_col, merged.max_col + 1):
                data[row - 1][col - 1] = value

    return pd.DataFrame(data)


def detect_header_rows(df: pd.DataFrame, max_scan: int = 5) -> List[int]:
    """Heuristically detect header rows (supports multi-row headers)."""
    if df.empty:
        return []

    header_rows: List[int] = []
    scan_limit = min(max_scan, len(df))

    for idx in range(scan_limit):
        row = df.iloc[idx]
        non_null = row.notna().sum()
        if non_null == 0:
            continue

        str_ratio = row.apply(lambda v: isinstance(v, str)).sum() / non_null
        if str_ratio >= 0.5:
            if header_rows and idx - header_rows[-1] > 1:
                break
            header_rows.append(idx)
        elif header_rows:
            break

    if not header_rows:
        header_rows = [0]

    return header_rows


def build_columns(df: pd.DataFrame, header_rows: List[int]) -> List[str]:
    """Flatten (potentially multi-row) headers into single-line column names."""
    if df.empty:
        return []

    header_df = df.iloc[header_rows].copy()
    header_df = header_df.ffill(axis=1)
    header_df = header_df.fillna("")
    header_df = header_df.apply(lambda col: col.map(lambda v: str(v).strip()))

    columns: List[str] = []
    for col_vals in zip(*header_df.values):
        parts = [part for part in col_vals if part]
        col_name = " / ".join(parts) if parts else f"col_{len(columns) + 1}"
        columns.append(col_name)

    return columns


def normalize_sheet(ws, sheet_name: str, sample_size: int) -> Tuple[dict, dict]:
    df = worksheet_to_dataframe(ws)
    merged_ranges = [str(rng) for rng in ws.merged_cells.ranges]
    header_rows = detect_header_rows(df)
    columns = build_columns(df, header_rows)

    data_df = df.drop(index=header_rows)
    data_df = data_df[data_df.notna().any(axis=1)]
    data_df.columns = columns if columns else list(range(len(data_df.columns)))

    normalized_rows = []
    for excel_idx, row in data_df.iterrows():
        record = {}
        for col, val in row.items():
            if pd.isna(val):
                continue
            record[col] = val
        record["_row_number"] = excel_idx + 1  # Excel rows are 1-based
        normalized_rows.append(record)

    sample_rows = normalized_rows[:sample_size]

    sheet_profile = {
        "sheet": sheet_name,
        "shape": (len(df.index), len(df.columns)),
        "header_rows": header_rows,
        "columns": columns,
        "merged_cells": merged_ranges,
        "sample_rows": sample_rows,
    }

    normalized_sheet = {
        "name": sheet_name,
        "header_rows": header_rows,
        "columns": columns,
        "merged_cells": merged_ranges,
        "rows": normalized_rows,
    }

    return sheet_profile, normalized_sheet


def main() -> None:
    parser = argparse.ArgumentParser(description="Parse Excel and normalize sheets.")
    parser.add_argument(
        "--excel",
        default="docs/data/TaskBreakDown_Bidding.xlsx",
        help="Path to Excel file.",
    )
    parser.add_argument(
        "--output",
        default="docs/product/_raw/excel-normalized.json",
        help="Path to write normalized JSON.",
    )
    parser.add_argument(
        "--sample-size",
        type=int,
        default=8,
        help="Number of sample rows to print per sheet.",
    )
    args = parser.parse_args()

    excel_path = Path(args.excel)
    output_path = Path(args.output)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = load_workbook(excel_path, data_only=True)

    profiles = []
    normalized = {
        "source": str(excel_path),
        "sheets": [],
    }

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        profile, normalized_sheet = normalize_sheet(ws, sheet_name, args.sample_size)
        profiles.append(profile)
        normalized["sheets"].append(normalized_sheet)

    print(f"Loaded workbook: {excel_path}")
    for profile in profiles:
        print(f"\nSheet: {profile['sheet']}")
        rows, cols = profile["shape"]
        print(f"  Shape: {rows} rows x {cols} cols")
        print(f"  Header rows: {profile['header_rows']}")
        print(f"  Merged cells: {len(profile['merged_cells'])}")
        if profile["merged_cells"][:5]:
            print(f"    Sample merges: {profile['merged_cells'][:5]}")
        print(f"  Columns: {profile['columns']}")
        print(f"  Sample rows (up to {len(profile['sample_rows'])}):")
        for sample in profile["sample_rows"]:
            print(f"    r{sample.get('_row_number')}: {sample}")

    with output_path.open("w", encoding="utf-8") as f:
        json.dump(normalized, f, ensure_ascii=False, indent=2, default=str)

    print(f"\nNormalized JSON written to {output_path}")


if __name__ == "__main__":
    main()
